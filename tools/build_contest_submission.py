from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement as SharedOxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


ROOT = Path(r"C:\Users\17480\Documents\琐碎小事儿\personal-cloud-drive")
OUT = ROOT / "output" / "contest-submission"
ASSETS = ROOT / "public"
DESIGN = ROOT / "design"
SCREENSHOTS = ROOT / "output" / "contest-submission" / "screenshots"

PROJECT = "智云引擎：面向高校科研资料管理的AI智能云网盘平台"
PLAN_TITLE = "《智云引擎AI智能云网盘平台商业计划书》"
PLAN_COVER_TITLE = "《智云引擎AI智能云网盘平台\n商业计划书》"
SCHOOL = "长江师范学院"
TEAM = "智云引擎队"
LEADER = "魏代平"
PHONE = "18580829563"
EMAIL = "wdp2467@163.com"


def set_cell_text(cell, text, font="楷体", size=12, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    if "@" in text or text.startswith("http"):
        font = "Calibri"
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_paragraph_font(paragraph, font="楷体", size=14, bold=False):
    for run in paragraph.runs:
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        run.font.size = Pt(size)
        run.bold = bold


def add_field(paragraph, field_code: str):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fld_end)


def setup_doc_styles(doc: Document):
    sec = doc.sections[0]
    setup_section_page(sec)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "楷体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(22)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = styles["Heading 1"]
    h1.font.name = "黑体"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(18)
    h1.font.bold = False
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.space_before = Pt(11)
    h1.paragraph_format.space_after = Pt(11)

    h2 = styles["Heading 2"]
    h2.font.name = "楷体"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    h2.font.size = Pt(16)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(11)
    h2.paragraph_format.space_after = Pt(11)

    h3 = styles["Heading 3"]
    h3.font.name = "楷体"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    h3.font.size = Pt(14)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("TOC 1", "TOC 2", "TOC 3", "toc 1", "toc 2", "toc 3"):
        if name in styles:
            toc_style = styles[name]
            toc_style.font.name = "楷体"
            toc_style._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
            toc_style.font.size = Pt(14)
            toc_style.font.color.rgb = RGBColor(0, 0, 0)
            toc_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            toc_style.paragraph_format.line_spacing = Pt(22)


def setup_section_page(sec):
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.0)


def clear_footer(section):
    section.footer.is_linked_to_previous = False
    for paragraph in section.footer.paragraphs:
        paragraph.clear()


def restart_page_number(section, start=1):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE")
    set_paragraph_font(p, "楷体", 10)


def add_para(doc, text="", style=None, first_line=True):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(22)
    if first_line and style is None:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    r.font.name = "楷体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    r.font.size = Pt(14)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(22)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        r = p.add_run(f"（{item[0]}）{item[1]}")
        r.font.name = "楷体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
        r.font.size = Pt(14)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = SharedOxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = SharedOxmlElement(tag)
            tc_borders.append(element)
        if edge_data is None:
            element.set(qn("w:val"), "nil")
            element.set(qn("w:sz"), "0")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "FFFFFF")
        else:
            element.set(qn("w:val"), edge_data.get("val", "single"))
            element.set(qn("w:sz"), str(edge_data.get("sz", 8)))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), edge_data.get("color", "000000"))


def apply_three_line_table(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
    line = {"val": "single", "sz": 8, "color": "000000"}
    bold_line = {"val": "single", "sz": 12, "color": "000000"}
    for cell in table.rows[0].cells:
        set_cell_border(cell, top=bold_line, bottom=line)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=bold_line)


def add_table(doc, title, headers, rows, widths):
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(title)
    run.font.name = "楷体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0, 0, 0)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (h, w) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, "楷体", 11, True)
        cell.width = Cm(w)
    for row in rows:
        cells = table.add_row().cells
        for i, (text, w) in enumerate(zip(row, widths)):
            cell = cells[i]
            set_cell_text(cell, text, "楷体", 10.5, False)
            cell.width = Cm(w)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER
    apply_three_line_table(table)
    doc.add_paragraph()
    return table


def add_figure(doc, image_name, caption, width_cm=14.2, page_break_before=False):
    image_path = SCREENSHOTS / image_name
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    p.add_run().add_break(WD_BREAK.LINE)
    r = p.add_run(caption)
    r.font.name = "楷体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0, 0, 0)


def build_business_plan():
    doc = Document()
    setup_doc_styles(doc)
    clear_footer(doc.sections[0])

    # Cover.
    logo = ASSETS / "dp-logo.png"
    if logo.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p_logo.paragraph_format.space_before = Pt(42)
        p_logo.add_run().add_picture(str(logo), width=Cm(3.0))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PLAN_COVER_TITLE)
    r.font.name = "楷体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    r.font.size = Pt(26)
    r.bold = False
    p.paragraph_format.space_after = Pt(42)

    info = [
        ("学校", SCHOOL),
        ("团队名称", TEAM),
        ("负责人", LEADER),
        ("联系方式", PHONE),
        ("邮箱", EMAIL),
        ("赛道", "AI创意赛道"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, (k, v) in enumerate(info):
        set_cell_text(table.rows[idx].cells[0], k, "楷体", 14, True)
        set_cell_text(table.rows[idx].cells[1], v, "楷体", 14)
        table.rows[idx].cells[0].width = Cm(4.2)
        table.rows[idx].cells[1].width = Cm(9.2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("重庆市首届高校AI大模型创新应用大赛")
    r.font.name = "楷体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    r.font.size = Pt(14)

    toc_section = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section_page(toc_section)
    clear_footer(toc_section)

    # TOC.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目  录")
    r.font.name = "楷体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    r.font.size = Pt(14)
    r.bold = True
    toc = doc.add_paragraph()
    add_field(toc, r'TOC \o "1-3" \h \z \u')
    set_paragraph_font(toc, "楷体", 14)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section_page(body_section)
    clear_footer(body_section)
    restart_page_number(body_section, 1)
    add_page_number(body_section)

    add_para(doc, PLAN_TITLE, None, False).alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(doc.paragraphs[-1], "楷体", 22, False)

    doc.add_heading("第一章 AI创意项目简介", level=1)
    add_para(doc, "智云引擎是一款面向高校学生、研究生、课题组与学院实验室的AI智能云网盘平台。项目以“安全存储、跨端访问、智能检索、资料管理”为核心，将传统网盘从单一文件存储工具升级为面向学习科研场景的资料管理助手。平台围绕高校科研资料体量大、格式杂、查找慢、共享难、权限弱等痛点，提供多用户独立空间、局域网与公网双通道访问、文件夹加密、批量上传下载、文档与图片预览、缓存加速、智能搜索和AI模式等功能，帮助用户在个人电脑或轻量服务器上构建可控、可扩展、低成本的个人云资料中心。")
    add_para(doc, "本项目选择教育与科研资料管理场景作为切入点。高校学生在课程学习、论文写作、竞赛申报、项目协作中会积累大量Word、PPT、PDF、图片、视频、压缩包和数据文件；课题组还会涉及过程文档、实验资料、会议记录、代码包和成果附件。传统本地文件夹依赖人工命名与层级记忆，普通网盘又往往存在公网访问不便、隐私控制不足、资料搜索能力弱和部署成本较高的问题。智云引擎通过AI大模型与智能搜索能力，将“存得下”进一步提升为“找得到、管得住、用得快”。")
    add_para(doc, "项目目前已完成可运行系统原型，具备多用户注册登录、账户隔离、文件上传下载、文件夹上传、批量下载、文件复制移动删除、文件夹密码保护、Office文档预览、图片浏览切换、局域网/公网访问状态识别、存储容量显示、缓存自动清理和智能搜索等功能。后续将围绕AI内容理解、自动分类、文档摘要、科研知识问答、资料关系图谱等方向持续迭代，形成面向高校科研场景的轻量化AI资料管理平台。")

    doc.add_heading("第二章 背景与问题分析", level=1)
    doc.add_heading("2.1 面向的问题与行业属性", level=2)
    add_para(doc, "本项目面向教育信息化与科研数据管理领域，聚焦高校个人与小型科研团队的资料管理问题。随着课程资料电子化、科研过程数字化、竞赛项目材料复杂化，学生和教师日常使用的文件种类越来越多，资料来源分散在个人电脑、手机、U盘、微信群、邮箱、云盘和实验室共享电脑中。用户经常遇到“文件在哪里”“哪个版本最新”“谁能查看”“手机照片如何批量备份”“公网和局域网如何稳定访问”等问题。")
    add_para(doc, "该问题具有明显的学习科研场景特征：资料以非结构化文件为主，文件类型多样，内容价值高但检索困难；资料涉及隐私、论文、项目申报、实验过程等内容，需要在共享效率和访问安全之间取得平衡；用户希望系统成本低、部署简单、访问灵活，并能在个人设备、宿舍网络、校园网和公网之间切换使用。")
    doc.add_heading("2.2 行业现状、问题、竞争与合作", level=2)
    add_para(doc, "当前资料管理工具主要包括操作系统文件夹、公共云盘、团队协作平台和学校内部系统。操作系统文件夹本地速度快，但跨端访问不便，检索能力弱；公共云盘使用门槛低，但用户对隐私、容量、会员费用和访问限制较敏感；团队协作平台适合企业流程，但对学生与小型课题组来说成本和学习门槛偏高；学校内部系统往往服务于教务或科研管理流程，难以覆盖个人资料日常管理。")
    add_para(doc, "智云引擎的差异化切入点不是与大型云盘正面竞争，而是面向高校个人和小型团队提供“轻量自部署、局域网优先、AI检索增强、权限清晰”的资料管理服务。它可以与已有云盘、学校系统和AI大模型平台形成互补：大型平台负责通用存储与公共服务，智云引擎负责用户可控的私有资料中心和科研资料组织。")
    doc.add_heading("2.3 行业趋势与发展方向", level=2)
    add_para(doc, "教育数字化和AI大模型应用正在推动资料管理方式变化。过去用户主要依赖目录层级和文件名记忆，现在更需要系统理解文件内容、识别文件类型、提供语义搜索、自动归类和知识问答。AI能力将使资料管理从“文件容器”走向“知识入口”，从被动存储走向主动辅助。")
    add_para(doc, "未来高校资料管理将呈现三个趋势：第一，私有化与轻量部署需求增强，用户希望重要资料可控可迁移；第二，多端协同与网络自适应成为基础能力，系统需要根据局域网和公网环境选择稳定访问路径；第三，AI内容理解成为核心增值能力，尤其在论文写作、竞赛材料整理、课题文档归档和历史资料复用中具有较高价值。")
    doc.add_heading("2.4 政策导向", level=2)
    add_para(doc, "国家和地方持续推动教育数字化、人工智能创新应用和数据安全治理。高校AI大模型创新应用大赛本身即体现了对AI技术落地应用的鼓励。智云引擎选择高校资料管理作为应用场景，符合AI赋能教育、提升科研效率、加强数据安全与隐私保护的方向。项目在设计上强调用户数据本地可控、账户隔离、文件夹加密和访问状态可见，有利于形成兼顾创新应用与安全治理的产品方案。")
    doc.add_heading("2.5 项目所处环境", level=2)
    add_para(doc, "项目处于从原型验证向产品化演进阶段。当前系统已在个人电脑作为服务器的环境下完成核心功能验证，支持局域网和公网访问，具备真实上传下载、预览、检索和权限控制能力。下一阶段将围绕高校学生与课题组的典型使用场景进行用户测试，沉淀需求反馈，并逐步补齐产品交互、部署工具、数据备份、AI索引能力和运维监控能力。")

    doc.add_heading("第三章 产品与技术", level=1)
    doc.add_heading("3.1 产品与技术的先进性和核心竞争力", level=2)
    add_para(doc, "智云引擎采用Web化云盘架构，用户可通过浏览器在电脑端、移动端、局域网和公网环境访问个人资料空间。平台在功能层面覆盖文件上传、文件夹上传、批量下载、文件复制、移动、删除、重命名、新建目录、多选操作、图片预览切换、Office文档预览、视频访问、缓存管理、存储空间统计等常用网盘能力。")
    add_para(doc, "平台的核心竞争力体现在四个方面：第一，面向高校场景的轻量化部署，用户可以用个人电脑或小型主机作为服务器，降低初期成本；第二，网络访问自适应，局域网环境优先提供更快访问，公网环境保证远程可用；第三，账户与文件夹权限隔离，多用户资料互不干扰，加密文件夹需验证密码后访问；第四，AI智能搜索和后续大模型能力可将资料管理从目录检索提升到内容理解。")
    add_table(
        doc,
        "表3-1 智云引擎已实现功能模块与用户价值",
        ["模块", "已实现能力", "对用户价值"],
        [
            ["账户系统", "注册登录、多用户空间隔离、密码找回", "保证不同用户资料独立，支持后续产品化"],
            ["文件管理", "上传、下载、文件夹上传、批量下载、复制、移动、删除、重命名", "覆盖日常资料管理高频操作"],
            ["安全权限", "文件夹密码、加密目录下载验证、登录密码验证", "保护论文、项目资料和隐私文件"],
            ["预览能力", "Word、PPT、PDF、图片、视频等预览与缓存", "减少反复下载，提高资料浏览效率"],
            ["AI能力", "智能搜索、AI模式入口、内容关键词检索", "降低海量资料定位成本"],
            ["网络能力", "局域网/公网访问状态显示、开机自启动与隧道守护", "提升远程访问稳定性"],
        ],
        [3.0, 6.0, 6.4],
    )
    add_figure(doc, "02-dashboard.png", "图3-1 智云引擎网盘主界面与访问状态展示")
    add_figure(doc, "03-smart-search.png", "图3-2 智能搜索结果与关键词高亮展示")
    add_para(doc, "在AI能力呈现上，平台已具备AI模式入口、全库问答界面和智能搜索联动能力。用户开启AI模式后，可以围绕当前网盘范围发起资料查找、结构整理和内容线索定位；系统会结合当前目录、文件名、路径和可提取文本形成资料管理上下文。现阶段该能力主要用于辅助定位资料和组织管理思路，后续将继续扩展至文档摘要、图片OCR检索、自动分类标签、课题资料问答和知识图谱等更深层的AI资料理解能力。")
    doc.add_page_break()
    add_figure(doc, "07-ai-mode.png", "图3-3 AI模式与全库资料问答界面展示", width_cm=12.6)
    add_figure(doc, "05-folder-password.png", "图3-4 加密文件夹访问验证弹窗展示")
    doc.add_heading("3.2 自主知识产权", level=2)
    add_para(doc, "项目代码和产品原型由团队自主设计与实现，已形成较完整的软件系统。项目后续可围绕软件著作权进行申报，保护系统在多用户网盘架构、文件夹密码管理、局域网/公网访问状态识别、AI资料搜索与缓存清理策略等方面的实现成果。若后续进一步形成自动分类、语义索引、科研资料知识图谱等算法模块，也可评估软件著作权、专有技术或相关专利布局。")
    doc.add_heading("3.3 研发力量与技术依托", level=2)
    add_para(doc, "团队负责人来自长江师范学院，具备人工智能相关专业背景，能够将AI大模型能力、Web开发、文件处理和产品设计结合到具体场景中。项目研发过程中已完成前后端交互、文件流处理、Office文档预览、多用户隔离、权限验证、公网隧道与局域网访问优化等多个工程模块。")
    add_para(doc, "技术依托主要包括Node.js后端服务、浏览器前端界面、Office文档解析与预览、OCR与文本提取、缓存管理、Cloudflare Tunnel公网访问、局域网服务发现和AI大模型接口能力。后续可结合学校实验室、指导教师资源和真实课题组用户进行场景验证。")
    doc.add_heading("3.4 国内与国际认证", level=2)
    add_para(doc, "项目当前处于软件原型和应用验证阶段，尚未进入正式认证流程。后续若面向学校、实验室或团队商业化部署，可根据实际服务范围完善软件著作权登记、信息系统安全测试、数据安全合规评估、用户隐私政策、网络安全等级保护适配等工作。")

    doc.add_heading("第四章 市场分析", level=1)
    doc.add_heading("4.1 发展前景", level=2)
    add_para(doc, "高校学生、研究生和课题组长期存在资料整理与共享需求。AI大模型普及后，用户不再满足于“云端存放文件”，而是希望系统能够理解文件内容，帮助用户快速检索、归档、摘要和复用资料。智云引擎处于教育数字化和AI资料管理的交汇点，具备从个人工具扩展到课题组轻协作平台的潜力。")
    doc.add_heading("4.2 定位清晰", level=2)
    add_para(doc, "项目初期定位为面向高校个人与小型团队的AI智能资料云盘。核心目标用户包括：课程资料较多的本科生和研究生，参与竞赛和项目申报的学生团队，需要管理论文、会议、实验与项目资料的课题组，以及希望部署轻量私有资料中心的教师或实验室。")
    doc.add_heading("4.3 目标市场", level=2)
    add_para(doc, "项目的初期目标市场以校内学生、课题组和竞赛团队为主，产品推广路径可从个人免费使用、课题组试点、学院实验室部署逐步展开。中期可拓展至高校创新创业团队、科研秘书、学生社团、培训机构等场景。远期可面向更广泛的小型团队，形成轻量私有云资料管理服务。")
    doc.add_heading("4.4 刚需痛点", level=2)
    add_bullets(
        doc,
        [
            ("1", "资料分散：手机照片、电脑文档、微信文件、U盘资料难以统一管理。"),
            ("2", "查找低效：文件名记不住、版本混乱、目录层级过深导致资料复用困难。"),
            ("3", "共享不稳：局域网快但地址变化，公网方便但速度与稳定性受网络影响。"),
            ("4", "隐私不足：论文、项目书、实验记录、账号资料等文件需要更细粒度保护。"),
            ("5", "工具割裂：存储、预览、搜索、加密、AI理解分别依赖不同工具，用户操作成本高。"),
        ],
    )
    doc.add_heading("4.5 商业模式", level=2)
    add_para(doc, "项目可采用“基础版免费、增强版付费、团队部署服务”的商业模式。基础版面向个人用户，提供核心存储与检索能力；增强版提供更大的预览缓存、AI内容索引、自动分类、批量资料整理、数据备份等能力；团队版面向课题组和实验室，提供多用户管理、专属部署、权限策略、技术支持和数据迁移服务。")
    add_table(
        doc,
        "表4-1 智云引擎商业模式分层设计",
        ["版本", "目标用户", "主要能力", "收费思路"],
        [
            ["个人基础版", "学生个人", "本地/公网访问、上传下载、预览、搜索", "免费或低价订阅"],
            ["个人增强版", "资料量较大的研究生", "AI索引、自动整理、扩展缓存、备份", "月/年订阅"],
            ["团队版", "课题组、竞赛队、实验室", "多账户空间、权限管理、部署支持", "按团队授权或部署服务收费"],
            ["学校合作版", "学院、实验室平台", "统一部署、账号体系适配、运维支持", "项目制服务"],
        ],
        [3.2, 3.6, 5.2, 3.4],
    )

    doc.add_heading("第五章 营销策略", level=1)
    doc.add_heading("5.1 传统组合策略", level=2)
    add_para(doc, "产品策略上，先以稳定可用的个人云网盘作为基础，突出AI资料搜索、文件夹加密、多端访问和课题组资料管理。价格策略上，初期以免费试用和低成本部署吸引校内用户，后续根据AI索引、存储规模和团队人数形成分级收费。渠道策略上，优先通过学院、实验室、竞赛社群、学生组织和指导教师资源进行试点。促销策略上，可围绕毕业论文资料整理、竞赛资料归档、课题组文档共享等具体场景制作案例。")
    doc.add_heading("5.2 网络组合策略", level=2)
    add_para(doc, "网络推广可通过项目官网、演示视频、校园公众号、B站/抖音产品演示、GitHub或开源社区展示等方式进行。重点不是泛泛宣传“云盘”，而是展示典型场景：手机照片一键备份、课题组资料加密共享、论文资料智能搜索、PPT/Word在线预览、局域网高速传输与公网远程访问。")
    doc.add_heading("5.3 新策略", level=2)
    add_para(doc, "项目可采用“场景任务驱动”的推广方式。例如面向研究生推出“论文资料整理包”，面向竞赛团队推出“项目材料归档包”，面向课题组推出“组会资料共享包”。通过具体任务引导用户体验AI搜索和资料管理能力，降低用户理解成本。")

    doc.add_heading("第六章 风险分析与控制", level=1)
    add_table(
        doc,
        "表6-1 项目主要风险与控制措施",
        ["风险类型", "主要表现", "控制措施"],
        [
            ["技术风险", "公网网络不稳定、文件预览兼容性差、大文件上传失败", "采用分片上传、缓存清理、状态监测、错误提示和多类型文件测试"],
            ["安全风险", "用户误共享、密码遗忘、目录权限混乱", "账户隔离、文件夹密码、操作确认、密码找回和权限校验"],
            ["市场风险", "用户习惯已有云盘，迁移动力不足", "从高校科研资料痛点切入，提供局域网速度、私有可控和AI搜索差异化"],
            ["运营风险", "团队运维能力不足，用户支持压力增加", "先做校内小范围试点，形成安装文档、故障提示和自动守护机制"],
            ["合规风险", "涉及用户隐私和资料安全", "明确隐私边界，默认用户自有服务器存储，减少不必要数据收集"],
            ["商业风险", "付费意愿不明确", "采用个人免费、团队付费、部署服务的渐进模式验证需求"],
        ],
        [2.8, 5.0, 7.0],
    )
    add_para(doc, "总体来看，项目的风险主要集中在产品稳定性、网络访问体验、AI能力深度和用户推广成本上。团队将以稳定为第一优先级，在保证上传下载、权限验证、数据隔离和缓存清理可靠的基础上逐步增强AI能力，避免过度承诺未成熟功能。")

    doc.add_heading("附件", level=1)
    add_bullets(
        doc,
        [
            ("1", "系统Demo：可通过局域网和公网访问的智云引擎个人云网盘原型。"),
            ("2", "功能截图：登录注册、多用户网盘、文件上传下载、智能搜索、文件夹加密、图片与文档预览。"),
            ("3", "后续材料：演示视频、用户测试反馈、软件著作权申请材料、项目源代码说明。"),
        ],
    )

    path = OUT / "智云引擎_AI创意赛道_商业计划书.docx"
    doc.save(path)
    return path


def rgb(hex_color: str):
    hex_color = hex_color.lstrip("#")
    return PptRGB(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def ppt_text(shape, text, size=24, bold=False, color="#0f172a", align=PP_ALIGN.LEFT):
    shape.text = text
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def wrap_cjk(text, max_chars=18):
    lines = []
    for part in str(text).split("\n"):
        current = ""
        for char in part:
            current += char
            if len(current) >= max_chars and char in "，。；、 ":
                lines.append(current.strip())
                current = ""
        while len(current) > max_chars:
            lines.append(current[:max_chars])
            current = current[max_chars:]
        if current.strip():
            lines.append(current.strip())
    return "\n".join(lines)


def add_title(slide, title, subtitle=None, dark=False):
    color = "#ffffff" if dark else "#0f172a"
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.42), PptInches(12.2), PptInches(0.55))
    ppt_text(box, title, 30, True, color)
    if subtitle:
        sub = slide.shapes.add_textbox(PptInches(0.58), PptInches(0.98), PptInches(10.8), PptInches(0.35))
        ppt_text(sub, subtitle, 14, False, "#b9d6f8" if dark else "#475569")


def add_footer(slide, index):
    box = slide.shapes.add_textbox(PptInches(11.55), PptInches(7.02), PptInches(1.1), PptInches(0.25))
    ppt_text(box, f"{index:02d}", 10, False, "#94a3b8")


def add_card(slide, x, y, w, h, title, body, fill="#ffffff", line="#d8e4f2", title_color="#0f172a", body_color="#475569"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = PptPt(1)
    t = slide.shapes.add_textbox(PptInches(x + 0.18), PptInches(y + 0.18), PptInches(w - 0.36), PptInches(0.35))
    ppt_text(t, title, 17, True, title_color)
    b = slide.shapes.add_textbox(PptInches(x + 0.18), PptInches(y + 0.62), PptInches(w - 0.36), PptInches(h - 0.72))
    b.text_frame.word_wrap = True
    ppt_text(b, wrap_cjk(body, max(12, int(w * 7.2))), 12.2, False, body_color)


def add_picture_frame(slide, image_name, x, y, w, h, caption=None):
    image_path = SCREENSHOTS / image_name
    if not image_path.exists():
        return
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame.fill.solid()
    frame.fill.fore_color.rgb = rgb("#ffffff")
    frame.line.color.rgb = rgb("#d8e4f2")
    frame.line.width = PptPt(1)
    slide.shapes.add_picture(str(image_path), PptInches(x + 0.08), PptInches(y + 0.08), width=PptInches(w - 0.16), height=PptInches(h - 0.36 if caption else h - 0.16))
    if caption:
        c = slide.shapes.add_textbox(PptInches(x + 0.15), PptInches(y + h - 0.25), PptInches(w - 0.3), PptInches(0.2))
        ppt_text(c, caption, 9.5, False, "#64748b", PP_ALIGN.CENTER)


def build_ppt():
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]
    logo = ASSETS / "dp-logo.png"
    preview = DESIGN / "ai-login-background-preview.png"

    def bg(slide, color="#f4f8fc"):
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        rect.fill.solid()
        rect.fill.fore_color.rgb = rgb(color)
        rect.line.fill.background()

    slides = []

    s = prs.slides.add_slide(blank)
    bg(s, "#071427")
    if preview.exists():
        s.shapes.add_picture(str(preview), 0, 0, width=prs.slide_width, height=prs.slide_height)
    shade = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(6.2), prs.slide_height)
    shade.fill.solid()
    shade.fill.fore_color.rgb = rgb("#071427")
    shade.fill.transparency = 12
    shade.line.fill.background()
    if logo.exists():
        s.shapes.add_picture(str(logo), PptInches(0.65), PptInches(0.56), width=PptInches(0.55))
    title = s.shapes.add_textbox(PptInches(0.72), PptInches(1.85), PptInches(5.8), PptInches(1.5))
    ppt_text(title, "智云引擎", 48, True, "#ffffff")
    sub = s.shapes.add_textbox(PptInches(0.76), PptInches(3.15), PptInches(5.9), PptInches(1.0))
    ppt_text(sub, "面向高校科研资料管理的\nAI智能云网盘平台", 24, False, "#60d4ff")
    meta = s.shapes.add_textbox(PptInches(0.78), PptInches(6.35), PptInches(5.5), PptInches(0.5))
    ppt_text(meta, f"{SCHOOL}｜{TEAM}｜负责人：{LEADER}", 13, False, "#cbd5e1")
    slides.append(s)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "高校科研资料管理正在从“存文件”变成“管知识”", "AI大模型让资料检索、归档和复用成为新的效率入口")
    add_card(s, 0.75, 1.55, 3.55, 4.55, "资料越来越多", "课程资料、论文文献、组会PPT、实验记录、竞赛材料、图片和视频持续增长。")
    add_card(s, 4.85, 1.55, 3.55, 4.55, "查找越来越慢", "目录层级、文件名和版本号依赖人工记忆，时间越久越难复用。")
    add_card(s, 8.95, 1.55, 3.55, 4.55, "安全越来越重要", "课题资料、项目申报和个人文件需要更清晰的权限隔离与加密保护。")
    add_footer(s, 2)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "智云引擎提供一个可控、可访问、可理解的个人资料中心")
    add_card(s, 0.75, 1.55, 3.7, 4.6, "可控", "用户可用个人电脑或小型主机部署，资料保存在自有存储空间，多用户账户相互隔离。", "#f8fbff")
    add_card(s, 4.82, 1.55, 3.7, 4.6, "可访问", "局域网高速访问，公网远程访问兜底，电脑端和移动端保持一致体验。", "#f8fbff")
    add_card(s, 8.9, 1.55, 3.7, 4.6, "可理解", "AI搜索和内容索引帮助用户从海量资料中快速定位真正需要的文件。", "#f8fbff")
    add_footer(s, 3)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "产品已经具备完整网盘原型，不停留在概念阶段")
    add_picture_frame(s, "02-dashboard.png", 0.7, 1.35, 7.45, 4.85, "主界面：账户隔离、访问状态、存储空间与文件列表")
    add_card(s, 8.55, 1.38, 3.75, 1.08, "文件管理", "上传、下载、文件夹上传、复制、移动、删除、批量下载")
    add_card(s, 8.55, 2.66, 3.75, 1.08, "权限安全", "注册登录、多用户隔离、文件夹密码、加密下载验证")
    add_card(s, 8.55, 3.94, 3.75, 1.08, "预览体验", "Word、PPT、PDF、图片、视频预览，缓存自动清理")
    add_card(s, 8.55, 5.22, 3.75, 1.08, "网络访问", "局域网/公网访问状态显示，开机自启动和隧道守护")
    add_footer(s, 4)

    s = prs.slides.add_slide(blank)
    bg(s, "#071427")
    add_title(s, "AI能力不是装饰，而是资料管理效率的核心增量", "从文件名检索走向内容理解、语义定位和知识复用", dark=True)
    add_picture_frame(s, "03-smart-search.png", 0.72, 1.5, 7.2, 4.82, "智能搜索：关键词命中、结果分类与路径定位")
    add_card(s, 8.3, 1.55, 3.95, 1.3, "当前已实现", "智能搜索、AI模式入口、内容关键词检索、多类型文件预览辅助", "#0f2a45", "#255d8a", "#e0f2fe", "#cbd5e1")
    add_card(s, 8.3, 3.15, 3.95, 1.3, "近期增强", "文档自动摘要、智能分类标签、OCR图片检索、科研资料归档建议", "#0f2a45", "#255d8a", "#e0f2fe", "#cbd5e1")
    add_card(s, 8.3, 4.75, 3.95, 1.3, "远期拓展", "课题资料问答、论文资料关系图谱、项目材料自动汇编、团队知识库", "#0f2a45", "#255d8a", "#e0f2fe", "#cbd5e1")
    add_footer(s, 5)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "核心用户从个人学生切入，逐步扩展到课题组和学院实验室")
    add_card(s, 0.75, 1.5, 3.55, 4.7, "学生个人", "手机照片备份、课程资料整理、论文资料检索、竞赛材料归档。")
    add_card(s, 4.85, 1.5, 3.55, 4.7, "课题组", "组会资料共享、实验文件归档、成员权限区分、历史资料复用。")
    add_card(s, 8.95, 1.5, 3.55, 4.7, "学院实验室", "轻量私有部署、统一资料空间、项目资料沉淀、运维可控。")
    add_footer(s, 6)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "商业模式采用个人免费试用、团队付费部署的渐进路径")
    tiers = [
        ("个人基础版", "免费/低价", "基础存储、预览、搜索和公网访问"),
        ("个人增强版", "订阅制", "AI索引、自动整理、扩展缓存和备份"),
        ("团队版", "授权/服务费", "多账户空间、权限策略、部署支持"),
        ("学校合作版", "项目制", "统一部署、账号适配、运维支持"),
    ]
    for i, (name, price, body) in enumerate(tiers):
        add_card(s, 0.75 + i * 3.05, 1.65, 2.65, 4.3, name, f"{price}\n\n{body}", "#f8fbff")
    add_footer(s, 7)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "差异化来自“私有可控 + 局域网优先 + AI资料管理”的组合")
    add_card(s, 0.82, 1.55, 3.6, 4.65, "相对公共云盘", "更强调自有存储、私密资料保护和高校科研场景，而不是通用文件托管。")
    add_card(s, 4.85, 1.55, 3.6, 4.65, "相对本地文件夹", "提供多端访问、权限管理、预览缓存和智能搜索，不再依赖人工目录记忆。")
    add_card(s, 8.88, 1.55, 3.6, 4.65, "相对团队协作平台", "部署更轻、成本更低、个人和小团队更容易上手。")
    add_footer(s, 8)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "系统架构围绕稳定访问、文件处理和AI检索三层展开")
    add_card(s, 0.7, 1.45, 3.2, 4.8, "访问层", "浏览器端\n移动端\n局域网访问\n公网隧道访问")
    add_card(s, 5.05, 1.45, 3.2, 4.8, "服务层", "账户认证\n文件管理\n权限校验\n缓存清理\n状态监测")
    add_card(s, 9.35, 1.45, 3.2, 4.8, "智能层", "内容提取\n智能搜索\nAI模式\n后续大模型问答")
    add_footer(s, 9)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "项目推进以稳定为第一优先级，逐步增强AI深度")
    milestones = [
        ("已完成", "可运行网盘原型，多用户隔离、上传下载、预览、加密、智能搜索"),
        ("近期", "完善比赛演示视频，优化AI搜索体验，补充用户测试反馈"),
        ("中期", "自动分类、摘要、OCR检索、资料归档建议"),
        ("远期", "课题组知识库、资料问答、学校/实验室部署服务"),
    ]
    for i, (a, b) in enumerate(milestones):
        add_card(s, 0.9 + i * 3.05, 1.8, 2.6, 3.9, a, b, "#ffffff")
    add_footer(s, 10)

    s = prs.slides.add_slide(blank)
    bg(s)
    add_title(s, "主要风险可通过工程稳定性和渐进式商业验证控制")
    add_picture_frame(s, "05-folder-password.png", 0.75, 1.4, 6.9, 4.9, "加密文件夹：进入目录前验证当前文件夹密码")
    add_card(s, 8.05, 1.55, 4.2, 1.25, "技术风险", "通过分片上传、缓存清理、状态监测和多文件类型测试提升稳定性。")
    add_card(s, 8.05, 3.2, 4.2, 1.25, "市场风险", "先从校内真实资料管理场景切入，验证学生和课题组需求。")
    add_card(s, 8.05, 4.85, 4.2, 1.25, "安全风险", "默认自有存储、账户隔离、文件夹密码和操作确认，减少误操作。")
    add_footer(s, 11)

    s = prs.slides.add_slide(blank)
    bg(s, "#071427")
    add_title(s, "智云引擎把普通网盘升级为高校科研资料的AI管理入口", "下一步将完成演示视频、用户试点和AI资料理解能力增强", dark=True)
    add_card(s, 1.0, 1.95, 3.3, 3.8, "可落地", "已有可运行系统原型，具备真实上传下载和权限控制能力。", "#0f2a45", "#255d8a", "#e0f2fe", "#cbd5e1")
    add_card(s, 5.0, 1.95, 3.3, 3.8, "有场景", "聚焦高校科研资料管理，痛点具体且使用频率高。", "#0f2a45", "#255d8a", "#e0f2fe", "#cbd5e1")
    add_card(s, 9.0, 1.95, 3.3, 3.8, "能扩展", "AI搜索、自动分类、摘要和资料问答可持续增强。", "#0f2a45", "#255d8a", "#e0f2fe", "#cbd5e1")
    add_footer(s, 12)

    path = OUT / "智云引擎_AI创意赛道_网评PPT.pptx"
    prs.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    docx_path = build_business_plan()
    pptx_path = build_ppt()
    print(docx_path)
    print(pptx_path)


if __name__ == "__main__":
    main()
