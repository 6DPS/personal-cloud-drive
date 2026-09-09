from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "软著合作开发协议_最终打印版.docx"

FONT_CN = "宋体"
FONT_EN = "Times New Roman"
TITLE = "计算机软件合作开发协议"

SOFTWARE_NAME = "DPSir智能云网盘服务软件"
SOFTWARE_VERSION = "V2.0"
SOFTWARE_FULL = f"{SOFTWARE_NAME} {SOFTWARE_VERSION}"


def set_run_font(run, size=12, bold=False):
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_format(paragraph, before=0, after=6, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.25
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_paragraph(doc, text="", *, size=12, bold=False, align=None, before=0, after=6, first_line=True):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    set_paragraph_format(paragraph, before=before, after=after, first_line=first_line)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_run_font(run, size=10)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(12)

    add_page_number(section.footer.paragraphs[0])

    add_paragraph(
        doc,
        TITLE,
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
        first_line=False,
    )

    for line in [
        "甲方：____________________________",
        "身份证号：____________________________",
        "",
        "乙方：____________________________",
        "身份证号：____________________________",
    ]:
        add_paragraph(doc, line, first_line=False, after=4 if line else 8)

    intro = (
        f'鉴于甲乙双方拟共同开发“{SOFTWARE_FULL}”（以下简称“本软件”），根据'
        "《中华人民共和国著作权法》《计算机软件保护条例》及其他相关法律法规，"
        "本着平等、自愿、诚实信用原则，经友好协商，就本软件合作开发及软件著作权"
        "归属等事项达成如下协议，以资共同遵守。"
    )
    add_paragraph(doc, intro, after=8)

    sections = [
        (
            "第一条 软件基本信息",
            [
                f"1. 软件名称：{SOFTWARE_NAME}。",
                f"2. 版本号：{SOFTWARE_VERSION}。",
                "3. 软件类型：应用软件。",
                "4. 软件用途：本软件用于个人和团队资料的云端存储、上传下载、文件管理、文件搜索、文件预览、权限控制、局域网/公网访问，以及基于人工智能大模型的全库问答、单文件对话、文档理解、内容辅助分析和图片文字识别等功能。",
            ],
        ),
        (
            "第二条 合作开发事项",
            [
                "1. 甲乙双方确认以合作开发方式共同完成本软件的需求分析、系统设计、程序开发、测试优化、人工智能能力接入、文档整理及成果申报工作。",
                "2. 双方同意在本软件开发完成后，以共同著作权人的身份申请计算机软件著作权登记。",
            ],
        ),
        (
            "第三条 双方分工",
            [
                "1. 甲方负责：需求分析、系统架构设计、前后端程序编写、文件管理与传输功能实现、账户与权限模块实现、云端访问能力实现、人工智能功能接口集成、系统联调、部署测试、源程序整理及申请材料初稿编制等工作。",
                "2. 乙方负责：技术路线指导、系统方案论证、人工智能功能设计讨论、开发过程指导、测试验收、性能与功能完善建议、软件说明文档审核及申报指导等工作。",
                "3. 双方确认，前述分工均对本软件开发完成具有实质性贡献。",
            ],
        ),
        (
            "第四条 开发起止时间",
            [
                "1. 合作开发开始日期：______年______月______日。",
                "2. 软件开发完成日期：______年______月______日。",
                "3. 双方确认，本协议签署日期不晚于软件开发完成日期。",
            ],
        ),
        (
            "第五条 著作权归属",
            [
                "1. 双方一致确认，本软件系甲乙双方合作开发完成的软件作品。",
                "2. 本软件的计算机软件著作权由甲乙双方共同享有。",
                "3. 双方同意在中国版权保护中心办理本软件计算机软件著作权登记，并在登记材料中列明双方为著作权人。",
            ],
        ),
        (
            "第六条 权利行使与成果使用",
            [
                "1. 双方可将本软件用于学习、科研、课程结题、竞赛、项目申报、毕业材料提交、职称或考核材料提交等非商业用途。",
                "2. 涉及本软件的对外转让、独占许可、排他许可、对外商业授权、对外投资、技术作价入股等重大处分事项，应事先取得双方书面同意。",
                "3. 双方对本软件进行升级、扩展、版本迭代、人工智能能力增强或衍生开发时，应就新增成果权属另行协商；未另行约定的，新增成果由实际完成方依法享有相应权利。",
            ],
        ),
        (
            "第七条 收益分配",
            [
                "1. 如本软件仅用于学习、科研、课程结题、项目申报、毕业材料提交及软件著作权登记，双方确认暂不涉及商业收益分配；如后续发生商业收益，双方另行协商确定。",
            ],
        ),
        (
            "第八条 署名方式",
            [
                "1. 双方同意，本软件在软件著作权登记、项目申报、说明文档及相关成果材料中的署名方式如下：",
                "   甲方：____________________________。",
                "   乙方：____________________________。",
                "2. 如双方另有署名顺序要求，可在提交前以书面方式确认。",
            ],
        ),
        (
            "第九条 保密条款",
            [
                "1. 未经对方书面同意，任何一方不得擅自向第三方披露本软件未公开的源代码、技术方案、模型接入配置、接口密钥、测试数据、文档内容及其他保密信息。",
                "2. 法律法规、司法机关、行政主管部门或软件著作权登记主管机构要求提供材料的，不受前款限制。",
            ],
        ),
        (
            "第十条 违约责任",
            [
                "1. 任一方违反本协议约定，导致对方遭受损失的，应承担相应赔偿责任。",
                "2. 任一方未经另一方同意，擅自处分本软件共有著作权或作出与本协议约定不一致的登记、授权、转让行为的，应承担由此产生的法律责任。",
            ],
        ),
        (
            "第十一条 争议解决",
            [
                "1. 因履行本协议发生争议，双方应先友好协商解决。",
                "2. 协商不成的，任一方可向有管辖权的人民法院提起诉讼。",
            ],
        ),
        (
            "第十二条 其他",
            [
                "1. 本协议自双方签字之日起生效。",
                "2. 本协议一式______份，甲乙双方各执______份，具有同等法律效力。",
                "3. 本协议可作为计算机软件著作权登记的权属证明材料提交。",
            ],
        ),
    ]

    for section_title, items in sections:
        add_paragraph(doc, section_title, bold=True, before=10, after=6, first_line=False)
        for item in items:
            add_paragraph(doc, item)

    doc.add_page_break()
    add_paragraph(
        doc,
        "签署页",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=20,
        first_line=False,
    )
    for line in [
        "甲方（签字）：____________________________",
        "签署日期：______年______月______日",
        "",
        "",
        "乙方（签字）：____________________________",
        "签署日期：______年______月______日",
    ]:
        add_paragraph(doc, line, after=12 if line else 18, first_line=False)
    add_paragraph(
        doc,
        "注：本页为协议签署页，甲乙双方签字后与协议正文具有同等效力。",
        size=10,
        first_line=False,
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
